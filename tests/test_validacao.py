from datetime import date

from openpyxl import load_workbook

from src.demo import seed
from src.validacao import export


def test_export_has_robot_results_inputs_and_formulas(tmp_db, tmp_path):
    seed(tmp_db, today=date(2026, 10, 13))
    out = tmp_path / "v.xlsx"
    n = export(out, ["demo-a"], db_path=tmp_db)
    assert n == 27  # só critérios automatizados e aplicáveis da última coleta
    wb = load_workbook(out)
    assert wb.sheetnames[:3] == ["Instruções", "Métricas", "Avaliação"]
    ws = wb["Avaliação"]
    assert ws["G2"].value in {"atende", "parcial", "nao_atende", "inconclusivo"}
    assert ws["P2"].value.startswith("=IF(") and ws["L2"].value is None
    assert any(dv.sqref and "L2" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert wb["Métricas"]["B15"].value.startswith("=IFERROR(")  # kappa


def test_carry_over_reuses_human_answers_and_flags_changes(tmp_db, tmp_path):
    seed(tmp_db, today=date(2026, 10, 13))
    first = tmp_path / "antes.xlsx"
    export(first, ["demo-a"], db_path=tmp_db)
    wb = load_workbook(first)
    ws = wb["Avaliação"]
    ws["L2"], ws["O2"] = "nao_atende", "conferido no portal"
    ws["G3"] = "inconclusivo" if ws["G3"].value != "inconclusivo" else "atende"  # simula robô diferente antes
    wb.save(first)
    second = tmp_path / "depois.xlsx"
    export(second, ["demo-a"], db_path=tmp_db, previous=first)
    ws2 = load_workbook(second)["Avaliação"]
    assert ws2["L2"].value == "nao_atende" and ws2["O2"].value == "conferido no portal"
    assert ws2["G3"].comment is not None and "Antes:" in ws2["G3"].comment.text
